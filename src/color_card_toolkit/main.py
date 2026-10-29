from __future__ import annotations

import os
import sys

from PySide6.QtWidgets import QApplication

from color_card_toolkit.core.ocr_engine import RapidOcrEngine
from color_card_toolkit.core.resources import flat_template_path
from color_card_toolkit.ui.main_window import MainWindow


def main() -> int:
    smoke_result = _run_smoke_check()
    if smoke_result is not None:
        return smoke_result

    app = QApplication(sys.argv)
    app.setApplicationName("线下色卡采集工具集")
    window = MainWindow()
    window.show()
    return app.exec()


def _run_smoke_check() -> int | None:
    mode = os.environ.get("COLOR_CARD_TOOLKIT_SMOKE", "").strip().lower()
    if not mode:
        return None
    if mode == "ocr":
        from pathlib import Path
        from PIL import Image, ImageDraw, ImageFont

        image = Image.new("RGB", (420, 100), "white")
        font = ImageFont.truetype(str(Path(os.environ["WINDIR"]) / "Fonts/arial.ttf"), 48)
        ImageDraw.Draw(image).text((20, 20), "PU12345", fill="black", font=font)
        blocks = RapidOcrEngine(intra_op_num_threads=1, inter_op_num_threads=1).recognize_image_object(image)
        if "PU12345" not in "".join(block.text.replace(" ", "") for block in blocks):
            raise RuntimeError("打包后的OCR未识别出测试文字")
        return 0
    if mode == "template":
        from pathlib import Path
        import tempfile
        from docx import Document
        from openpyxl import Workbook
        from color_card_toolkit.core.spu_label_generator import convert_spu_excel_to_docx

        if not flat_template_path().exists():
            raise FileNotFoundError(flat_template_path())
        Document(str(flat_template_path()))
        with tempfile.TemporaryDirectory(prefix="color-card-template-check-") as folder:
            root = Path(folder)
            workbook = Workbook()
            workbook.active["A1"] = "测试标签"
            workbook.save(root / "sample.xlsx")
            output = convert_spu_excel_to_docx(root / "sample.xlsx", root / "sample.docx", start_row=1, start_column=1)
            if Document(str(output)).tables[0].cell(0, 0).text != "测试标签":
                raise RuntimeError("打包后的模板生成失败")
        return 0
    if mode == "ui":
        from PySide6.QtCore import Qt
        from PySide6.QtGui import QInputMethodEvent, QImageReader
        from PySide6.QtWidgets import QLineEdit

        app = QApplication([sys.argv[0]])
        window = MainWindow()
        window.setAttribute(Qt.WA_DontShowOnScreen)
        window.show()
        for show in (window.show_stack_to_flat, window.show_scan_rename, window.show_main_image_crop, window.show_spu_label):
            show()
            app.processEvents()
        entry = QLineEdit()
        event = QInputMethodEvent()
        event.setCommitString("主图中文输入")
        app.sendEvent(entry, event)
        if entry.text() != "主图中文输入":
            raise RuntimeError("中文输入事件未通过")
        formats = {bytes(value).decode() for value in QImageReader.supportedImageFormats()}
        if not {"jpeg", "png"}.issubset(formats):
            raise RuntimeError("缺少JPEG/PNG界面图片支持")
        window.close()
        return 0
    if mode == "main-ruler":
        import json
        from pathlib import Path
        import tempfile
        from PIL import Image
        from color_card_toolkit.core.ruler_inspection import inspect_image_ruler
        from color_card_toolkit.core.image_rename import crop_image_from_geometry

        source = Path(os.environ["COLOR_CARD_TOOLKIT_SMOKE_IMAGE"])
        inspection = inspect_image_ruler(source, crop_width_cm=5, crop_height_cm=5, search_ratio=0.20, suggest_retry=False)
        if not inspection.passed:
            raise RuntimeError(inspection.message)
        with tempfile.TemporaryDirectory(prefix="color-card-ruler-check-") as folder:
            output = Path(folder) / "main.jpg"
            crop_image_from_geometry(source, output, 5, 5, inspection.geometry)
            with Image.open(output) as image:
                marker = json.loads(image.getexif().get_ifd(34665)[37510][8:])
                if marker["width_cm"] != 5 or marker["height_cm"] != 5:
                    raise RuntimeError("打包后的标尺裁剪信息无效")
        return 0
    if mode == "main-marker":
        import io
        import json
        from PIL import Image
        from color_card_toolkit import __version__
        from color_card_toolkit.core.main_image_marker import marker_exif

        for image_format in ("JPEG", "PNG"):
            output = io.BytesIO()
            Image.new("RGB", (20, 20)).save(output, image_format, exif=marker_exif(5, 10, 2, 3))
            output.seek(0)
            with Image.open(output) as image:
                marker = json.loads(image.getexif().get_ifd(34665)[37510][8:])
            if marker["tool_version"] != __version__:
                raise RuntimeError("图片标记版本与软件版本不一致")
        return 0
    raise ValueError(
        f"Unknown smoke check mode: {mode}. "
        "The old local recognize smoke check was removed with the YOLO runtime."
    )


if __name__ == "__main__":
    raise SystemExit(main())
